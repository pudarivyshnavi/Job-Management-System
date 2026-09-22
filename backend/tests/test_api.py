"""Automated API tests for the Job Management System backend.

Run from the backend folder:
    .venv/Scripts/python -m pytest tests/ -v      (Windows)
    .venv/bin/python -m pytest tests/ -v          (Linux/macOS)
"""

import io
import json
import shutil
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app import services
from app import ai_service
from app import data_store
from app.main import app


@pytest.fixture()
def client(tmp_path):
    """Serve the app against isolated JSON collections and a temporary recruiter."""
    original_file = services.JOBS_FILE
    original_data_dir = data_store.DATA_DIR
    test_data_dir = tmp_path / "data"
    test_data_dir.mkdir()
    for source in original_data_dir.glob("*.json"):
        shutil.copy(source, test_data_dir / source.name)
    test_file = test_data_dir / "jobs.json"
    services.JOBS_FILE = test_file
    data_store.DATA_DIR = test_data_dir
    with TestClient(app, raise_server_exceptions=False) as test_client:
        signup = test_client.post(
            "/api/auth/signup",
            json={
                "name": "Test Recruiter",
                "email": "test-recruiter@example.test",
                "password": "testing-password",
                "confirm_password": "testing-password",
            },
        )
        assert signup.status_code == 201
        login = test_client.post(
            "/api/auth/login",
            json={
                "email": "test-recruiter@example.test",
                "password": "testing-password",
            },
        )
        test_client.headers["Authorization"] = f"Bearer {login.json()['token']}"
        yield test_client
    services.JOBS_FILE = original_file
    data_store.DATA_DIR = original_data_dir


VALID_JOB = {
    "title": "Automation Engineer",
    "company": "TestWorks",
    "location": "Remote",
    "description": "A fictional job used by automated tests.",
    "required_skills": ["Python", "Testing"],
    "experience": "2-4 years",
    "employment_type": "Full-time",
    "status": "Active",
}


# ---------------------------------------------------------------- read ----


def test_root_health(client):
    response = client.get("/")
    assert response.status_code == 200
    assert response.json() == {"status": "ok", "docs": "/docs"}


def test_health_endpoint(client):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "healthy"}


def test_auth_signup_login_me_and_duplicate_handling(client):
    duplicate = client.post(
        "/api/auth/signup",
        json={
            "name": "Another Recruiter",
            "email": "test-recruiter@example.test",
            "password": "testing-password",
            "confirm_password": "testing-password",
        },
    )
    assert duplicate.status_code == 409
    invalid = client.post(
        "/api/auth/login",
        json={"email": "test-recruiter@example.test", "password": "wrong-password"},
    )
    assert invalid.status_code == 401
    assert client.get("/api/auth/me").json()["role"] == "Recruiter"


def test_protected_resource_requires_authentication(client):
    token = client.headers.pop("Authorization")
    response = client.get("/api/candidates")
    assert response.status_code == 401
    client.headers["Authorization"] = token


def test_demo_resources_and_local_ai(client):
    assert client.get("/api/candidates").json()[0]["name"] == "Aarav Mehta"
    assert client.get("/api/applications").json()[0]["stage"] == "Shortlisted"
    assert client.get("/api/notifications").json()[0]["read"] is False
    analytics = client.get("/api/analytics").json()
    assert analytics["demo"] is True
    result = client.post("/api/ai/generate", json={"prompt": "Python platform work"})
    assert result.status_code == 200
    assert result.json()["source"] == "deterministic local assistant"


def test_list_jobs_returns_sample_data(client):
    response = client.get("/api/jobs")
    assert response.status_code == 200
    jobs = response.json()
    assert len(jobs) >= 15
    assert all(
        {"id", "title", "company", "location", "status", "created_date"} <= set(job)
        for job in jobs
    )


def test_get_one_job(client):
    response = client.get("/api/jobs/1")
    assert response.status_code == 200
    assert response.json()["id"] == 1


def test_get_missing_job_returns_404(client):
    response = client.get("/api/jobs/9999")
    assert response.status_code == 404
    assert "not found" in response.json()["detail"].lower()


# --------------------------------------------------------------- create ----


def test_create_job_returns_201_and_persists(client):
    response = client.post("/api/jobs", json=VALID_JOB)
    assert response.status_code == 201
    created = response.json()
    assert created["title"] == "Automation Engineer"
    assert created["created_date"]

    # The created job must survive a fresh read from disk.
    stored = json.loads(services.JOBS_FILE.read_text(encoding="utf-8"))
    assert any(job["title"] == "Automation Engineer" for job in stored)


def test_create_job_rejects_missing_fields(client):
    response = client.post("/api/jobs", json={"title": "Only a title"})
    assert response.status_code == 422


def test_create_job_rejects_whitespace_only_strings(client):
    payload = {**VALID_JOB, "company": "   "}
    response = client.post("/api/jobs", json=payload)
    assert response.status_code == 422


def test_create_job_rejects_invalid_enums(client):
    payload = {**VALID_JOB, "employment_type": "Shift-work"}
    response = client.post("/api/jobs", json=payload)
    assert response.status_code == 422

    payload = {**VALID_JOB, "status": "Closed"}
    response = client.post("/api/jobs", json=payload)
    assert response.status_code == 422


def test_create_job_rejects_empty_skills_list(client):
    payload = {**VALID_JOB, "required_skills": []}
    response = client.post("/api/jobs", json=payload)
    assert response.status_code == 422


# --------------------------------------------------------------- update ----


def test_update_job_keeps_id_and_created_date(client):
    before = client.get("/api/jobs/2").json()
    payload = {**VALID_JOB, "title": "Senior Data Analyst"}
    response = client.put("/api/jobs/2", json=payload)
    assert response.status_code == 200
    updated = response.json()
    assert updated["title"] == "Senior Data Analyst"
    assert updated["id"] == before["id"]
    assert updated["created_date"] == before["created_date"]


def test_update_missing_job_returns_404(client):
    response = client.put("/api/jobs/9999", json=VALID_JOB)
    assert response.status_code == 404


# -------------------------------------------------------------- archive ----


def test_archive_sets_status_and_keeps_record(client):
    total_before = len(client.get("/api/jobs").json())
    response = client.delete("/api/jobs/1")
    assert response.status_code == 200
    assert response.json()["status"] == "Archived"

    # The record is still on disk, not removed.
    stored = json.loads(services.JOBS_FILE.read_text(encoding="utf-8"))
    assert len(stored) == total_before
    assert any(job["id"] == 1 and job["status"] == "Archived" for job in stored)


def test_archived_job_found_under_status_filter(client):
    client.delete("/api/jobs/1")
    response = client.get("/api/jobs", params={"status": "Archived"})
    assert response.status_code == 200
    assert any(job["id"] == 1 for job in response.json())


def test_archive_missing_job_returns_404(client):
    response = client.delete("/api/jobs/9999")
    assert response.status_code == 404


# ------------------------------------------------------- search/filters ----

# The search tests below create their own probe jobs so they do not depend on
# the exact titles/companies in the sample data (which users may edit via the
# UI before running the suite).


def test_search_matches_title_partially_and_case_insensitively(client):
    probe = {**VALID_JOB, "title": "Zprobe Zearch Pyth Case"}
    created = client.post("/api/jobs", json=probe).json()
    response = client.get("/api/jobs", params={"search": "pyth"})
    titles = [job["title"] for job in response.json()]
    assert created["title"] in titles  # partial match, ignoring case


def test_search_matches_company(client):
    probe = {
        **VALID_JOB,
        "title": "Company Search Probe",
        "company": "Zimbus Probe Labs",
    }
    created = client.post("/api/jobs", json=probe).json()
    response = client.get("/api/jobs", params={"search": "zimbus"})
    assert any(job["id"] == created["id"] for job in response.json())


def test_search_with_no_match_returns_empty_list(client):
    response = client.get("/api/jobs", params={"search": "zzz-no-match"})
    assert response.status_code == 200
    assert response.json() == []


def test_filter_by_location(client):
    response = client.get("/api/jobs", params={"location": "Hyderabad"})
    assert all(job["location"].lower() == "hyderabad" for job in response.json())
    assert len(response.json()) > 0


def test_filter_by_experience(client):
    response = client.get("/api/jobs", params={"experience": "2-4 years"})
    assert all(job["experience"] == "2-4 years" for job in response.json())


def test_filter_by_employment_type(client):
    response = client.get("/api/jobs", params={"employment_type": "Internship"})
    assert all(job["employment_type"] == "Internship" for job in response.json())


def test_filter_by_status(client):
    response = client.get("/api/jobs", params={"status": "Draft"})
    assert all(job["status"] == "Draft" for job in response.json())


def test_combined_filters_and_search(client):
    probe = {
        **VALID_JOB,
        "title": "Combined Probe pyth",
        "location": "Hyderabad",
        "employment_type": "Full-time",
        "status": "Active",
    }
    created = client.post("/api/jobs", json=probe).json()
    response = client.get(
        "/api/jobs",
        params={
            "location": "Hyderabad",
            "employment_type": "Full-time",
            "status": "Active",
            "search": "pyth",
        },
    )
    jobs = response.json()
    assert any(job["id"] == created["id"] for job in jobs)
    assert all(
        job["location"] == "Hyderabad"
        and job["employment_type"] == "Full-time"
        and job["status"] == "Active"
        for job in jobs
    )


# ------------------------------------------------------------ data file ----


def test_corrupt_json_file_returns_500(client, tmp_path):
    services.JOBS_FILE.write_text("{not valid json", encoding="utf-8")
    response = client.get("/api/jobs")
    assert response.status_code == 500
    assert "invalid JSON" in response.json()["detail"]


def test_ai_unavailable_uses_explicit_local_fallback(client, monkeypatch):
    monkeypatch.setattr(
        ai_service,
        "ollama_generate",
        lambda prompt: ai_service.AiResult(
            False, "unavailable", message=ai_service.UNAVAILABLE_MESSAGE
        ),
    )
    response = client.post("/api/ai/generate", json={"prompt": "Python API engineer"})
    assert response.status_code == 200
    assert response.json()["available"] is False
    assert response.json()["source"] == "deterministic local assistant"
    assert "currently unavailable" in response.json()["message"]


def test_ai_service_errors_are_safe(monkeypatch):
    monkeypatch.setattr(
        ai_service.urllib.request,
        "urlopen",
        lambda *args, **kwargs: (_ for _ in ()).throw(TimeoutError()),
    )
    result = ai_service.ollama_generate("brief")
    assert result.available is False
    assert result.message == ai_service.UNAVAILABLE_MESSAGE


def test_advanced_search_and_job_lifecycle(client):
    response = client.get(
        "/api/jobs",
        params={"title": "Python", "skills": "FastAPI", "work_mode": "Hybrid"},
    )
    assert response.status_code == 200
    assert all("Python" in job["title"] for job in response.json())
    duplicate = client.post("/api/jobs/1/duplicate")
    assert duplicate.status_code == 201
    assert duplicate.json()["status"] == "Draft"
    closed = client.post("/api/jobs/1/close")
    assert closed.json()["workflow_status"] == "Closed"
    reopened = client.post("/api/jobs/1/reopen")
    assert reopened.json()["workflow_status"] == "Open"
    client.delete("/api/jobs/1")
    restored = client.post("/api/jobs/1/restore")
    assert restored.json()["status"] == "Active"


def test_matching_endpoint_is_transparent(client):
    response = client.get("/api/matching/1/1")
    assert response.status_code == 200
    payload = response.json()
    assert 0 <= payload["score"] <= 100
    assert "60% required skills" in payload["calculation"]
    assert "matching_skills" in payload and "missing_skills" in payload


def test_resume_txt_analysis_and_validation(client):
    resume = b"Mira Shah\nBSc Computer Science, Example University\nPython FastAPI React\n5 years experience"
    response = client.post(
        "/api/resumes/analyse", files={"file": ("resume.txt", resume, "text/plain")}
    )
    assert response.status_code == 200
    assert "Python" in response.json()["skills"]
    invalid = client.post(
        "/api/resumes/analyse",
        files={"file": ("resume.exe", b"bad", "application/octet-stream")},
    )
    assert invalid.status_code == 422


def test_application_and_interview_contracts(client):
    applications = client.get("/api/applications", params={"status": "Shortlisted"})
    assert applications.status_code == 200
    assert all(item["stage"] == "Shortlisted" for item in applications.json())
    interviews = client.get("/api/interviews")
    assert interviews.status_code == 200
    assert interviews.json()[0]["status"] in {"Scheduled", "Completed"}


def test_resume_analysis_accepts_valid_txt_pdf_and_docx(client):
    txt_resume = (
        b"Aarav Mehta\nPython FastAPI SQL\n5 years experience\nBSc Computer Science"
    )
    response = client.post(
        "/api/resumes/analyse",
        files={"file": ("resume.txt", txt_resume, "text/plain")},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["candidate"]["name"] == "Aarav Mehta"
    assert "Python" in payload["skills"]

    from docx import Document

    doc = Document()
    doc.add_paragraph("Mira Shah")
    doc.add_paragraph("React TypeScript testing")
    doc.add_paragraph("3 years experience")
    docx_bytes = io.BytesIO()
    doc.save(docx_bytes)
    docx_response = client.post(
        "/api/resumes/analyse",
        files={
            "file": (
                "resume.docx",
                docx_bytes.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert docx_response.status_code == 200
    assert "React" in docx_response.json()["skills"]

    pdf_payload = b"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>
endobj
4 0 obj
<< /Length 75 >>
stream
BT
/F1 18 Tf
50 80 Td
(Aarav Mehta Python FastAPI SQL) Tj
ET
endstream
endobj
5 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
xref
0 6
0000000000 65535 f
0000000010 00000 n
0000000063 00000 n
0000000125 00000 n
0000000377 00000 n
0000000840 00000 n
trailer
<< /Root 1 0 R /Size 6 >>
startxref
905
%%EOF
"""
    pdf_response = client.post(
        "/api/resumes/analyse",
        files={"file": ("resume.pdf", pdf_payload, "application/pdf")},
    )
    assert pdf_response.status_code == 200
    assert "summary" in pdf_response.json()


def test_resume_analysis_rejects_empty_oversized_and_invalid_uploads(client):
    empty_response = client.post(
        "/api/resumes/analyse",
        files={"file": ("empty.txt", b"", "text/plain")},
    )
    assert empty_response.status_code == 422
    assert "readable text" in empty_response.json()["detail"].lower()

    oversized = b"A" * (6 * 1024 * 1024)
    oversized_response = client.post(
        "/api/resumes/analyse",
        files={"file": ("oversized.txt", oversized, "text/plain")},
    )
    assert oversized_response.status_code == 422
    assert "5 mb" in oversized_response.json()["detail"].lower()

    invalid_response = client.post(
        "/api/resumes/analyse",
        files={"file": ("resume.exe", b"not a resume", "application/octet-stream")},
    )
    assert invalid_response.status_code == 422
    assert "pdf, docx, or txt" in invalid_response.json()["detail"].lower()
